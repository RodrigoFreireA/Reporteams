import os
import tempfile
import unittest

from openpyxl import Workbook

import pandas as pd

from generate_dashboard import build_roadmap_from_df, compute_json, load_base


class PlannerFormatTests(unittest.TestCase):
    def _save_workbook(self, wb):
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        return path

    def _new_format_workbook(self):
        wb = Workbook()
        ws_plan = wb.active
        ws_plan.title = "Plano"
        ws_plan.append(["ID do plano", "Nome do plano", "Data da exportação "])
        ws_plan.append(["plan-1", "Sprint Nova", "2026-05-14"])

        ws_tasks = wb.create_sheet("Dados Consolidados")
        ws_tasks.append(
            [
                "Identificação da tarefa",
                "Nome da tarefa",
                "Categoria",
                "Meta",
                "Status",
                "Prioridade",
                "Atribuído a",
                "Criado por",
                "Criado em",
                "Data de conclusão",
                "Data de início",
                "É Recorrente",
                "Atrasados",
                "Concluído em",
                "Concluída por",
                "Itens concluídos da lista de verificação",
                "Itens da lista de verificação",
                "Rótulos",
                "Notas",
            ]
        )
        ws_tasks.append(
            [
                "task-1",
                "Sprint Goal",
                "Concluído",
                "",
                "Concluída",
                "Média",
                "Gerente",
                "Gerente",
                "2026-05-01",
                "2026-05-10",
                "2026-05-01",
                "",
                "false",
                "2026-05-10",
                "Gerente",
                "",
                "",
                "FLUXO.CONTINUO;.GP",
                "Sprint goal: Entregar o formato novo.",
            ]
        )
        ws_tasks.append(
            [
                "task-2",
                "Implementar upload",
                "Em Andamento",
                "",
                "Em andamento",
                "Média",
                "Dev",
                "Gerente",
                "2026-05-02",
                "2026-05-14",
                "2026-05-02",
                "",
                "false",
                "",
                "",
                "1/3",
                "A;B;C",
                "HU001 - Login [3SP];.DEV",
                "",
            ]
        )
        return wb

    def _legacy_format_workbook(self, task_sheet_name="Tarefas"):
        wb = Workbook()
        ws = wb.active
        ws.title = task_sheet_name
        ws.append(
            [
                "Task name",
                "Bucket name",
                "Assigned to",
                "Status",
                "Date completed",
                "Start date",
                "Due date",
                "Labels",
                "Notes",
            ]
        )
        ws.append(
            [
                "Implementar legado",
                "Concluido",
                "Dev",
                "Concluida",
                "2026-05-08",
                "2026-05-01",
                "2026-05-08",
                "HU002 - Legado [2SP];.DEV",
                "",
            ]
        )
        ws_plan = wb.create_sheet("Nome do plano")
        ws_plan.append(["Nome do plano", "Sprint Antiga"])
        ws_plan.append(["Exportado em", "2026-05-09"])
        return wb

    def test_compute_json_reads_new_teams_format(self):
        path = self._save_workbook(self._new_format_workbook())
        try:
            data = compute_json(path, planner_format="teams_new")
        finally:
            os.unlink(path)

        self.assertEqual(data["meta"]["sprint_name"], "Sprint Nova")
        self.assertEqual(data["meta"]["export_date"], "2026-05-14")
        self.assertEqual(data["meta"]["planner_format"], "teams_new")
        self.assertEqual(data["meta"]["sprint_goal"], "Entregar o formato novo.")
        self.assertEqual(data["kpis"]["total"], 2)
        self.assertEqual(data["kpis"]["done"], 1)
        self.assertEqual(data["kpis"]["pending"], 1)
        self.assertEqual(data["kpis"]["hu_count"], 1)
        self.assertEqual(data["burndown_sp"]["total_sp"], 3)

    def test_compute_json_extracts_hu_and_sp_from_task_title(self):
        wb = self._legacy_format_workbook()
        ws = wb["Tarefas"]
        ws["A2"] = "[HU008 - Blablabla bla bla [13SP]] tarefa 002 bla blabla"
        ws["H2"] = ".DEV"

        path = self._save_workbook(wb)
        try:
            data = compute_json(path, planner_format="legacy")
        finally:
            os.unlink(path)

        self.assertEqual(data["hu_list"], [["HU008", 1, 1]])
        self.assertEqual(data["hu_full_names"], {"HU008": "HU008 - Blablabla bla bla [13SP]"})
        self.assertEqual(data["burndown_sp"]["hu_sp"], {"HU008": 13})
        self.assertEqual(data["burndown_sp"]["total_sp"], 13)
        self.assertEqual(data["kpis"]["sem_hu"], 0)
        self.assertEqual(data["kpis"]["hu_count"], 1)
        self.assertEqual(data["kpis"]["storypoints"], 13)
        self.assertEqual(data["task_rows"][0]["hu"], "HU008")

    def test_compute_json_prioritizes_title_hu_over_label_hu(self):
        wb = self._legacy_format_workbook()
        ws = wb["Tarefas"]
        ws["A2"] = "[HU008 - Titulo manda [13SP]] tarefa 002"
        ws["H2"] = "HU009 - Rotulo antigo [5SP];.DEV"

        path = self._save_workbook(wb)
        try:
            data = compute_json(path, planner_format="legacy")
        finally:
            os.unlink(path)

        self.assertEqual(data["hu_list"], [["HU008", 1, 1]])
        self.assertEqual(data["hu_full_names"], {"HU008": "HU008 - Titulo manda [13SP]"})
        self.assertEqual(data["burndown_sp"]["hu_sp"], {"HU008": 13})

    def test_compute_json_still_falls_back_to_hu_label(self):
        wb = self._legacy_format_workbook()
        path = self._save_workbook(wb)
        try:
            data = compute_json(path, planner_format="legacy")
        finally:
            os.unlink(path)

        self.assertEqual(data["hu_list"], [["HU002", 1, 1]])
        self.assertEqual(data["hu_full_names"], {"HU002": "HU002 - Legado [2SP]"})
        self.assertEqual(data["burndown_sp"]["hu_sp"], {"HU002": 2})

    def test_compute_json_uses_title_profile_marker_without_profile_label(self):
        wb = self._legacy_format_workbook()
        ws = wb["Tarefas"]
        ws["A2"] = "[.DEV] Implementar legado"
        ws["H2"] = "HU002 - Legado [2SP]"

        path = self._save_workbook(wb)
        try:
            data = compute_json(path, planner_format="legacy")
        finally:
            os.unlink(path)

        self.assertEqual(data["area_rows"], [[".DEV", 1, 0]])
        self.assertEqual(data["task_rows"][0]["profile_labels"], ".DEV")
        self.assertEqual(data["task_rows"][0]["profile_base"], "Devs")
        self.assertEqual(data["task_rows"][0]["wip_phase"], "Em Desenvolvimento")
        self.assertEqual(data["task_rows"][0]["hu"], "HU002")

    def test_compute_json_prioritizes_title_profile_marker_over_exported_label(self):
        wb = self._legacy_format_workbook()
        ws = wb["Tarefas"]
        ws["A2"] = "[ARQUITETURA] Ajustar desenho tecnico"
        ws["H2"] = "HU009 - Rotulo antigo [5SP];.DEV"

        path = self._save_workbook(wb)
        try:
            data = compute_json(path, planner_format="legacy")
        finally:
            os.unlink(path)

        self.assertEqual(data["area_rows"], [[".ARQUITETURA", 1, 0]])
        self.assertEqual(data["task_rows"][0]["profile_labels"], ".ARQUITETURA")
        self.assertEqual(data["task_rows"][0]["profile_base"], "Arquitetura")
        self.assertEqual(data["task_rows"][0]["hu"], "HU009")

    def test_roadmap_card_normalizes_unlabeled_goal_and_marco(self):
        df = pd.DataFrame(
            [
                {
                    "tarefa": "Roadmap",
                    "notas": (
                        "15/05/2026\n"
                        "Disponibilizar primeira versao\n"
                        "Publicacao do MVP\n"
                        "30/06/2026  Validar operacao  Encerrar piloto\n"
                        "15/07/2026\n"
                        "Goal: Escalar uso\n"
                        "Marco: Liberacao geral"
                    ),
                }
            ]
        )

        items = build_roadmap_from_df(df)

        self.assertEqual(
            items,
            [
                (
                    "15/05/2026\n"
                    "Goal: Disponibilizar primeira versao\n"
                    "Marco: Publicacao do MVP",
                    10,
                ),
                (
                    "30/06/2026\n"
                    "Goal: Validar operacao\n"
                    "Marco: Encerrar piloto",
                    -10,
                ),
                (
                    "15/07/2026\n"
                    "Goal: Escalar uso\n"
                    "Marco: Liberacao geral",
                    40,
                ),
            ],
        )

    def test_roadmap_accepts_named_task_and_iso_dates(self):
        df = pd.DataFrame(
            [
                {
                    "tarefa": "Roadmap do produto",
                    "notas": (
                        "2026-05-15\n"
                        "Disponibilizar primeira versao\n"
                        "Publicacao do MVP\n"
                        "15-06-2026\tValidar operacao\tEncerrar piloto"
                    ),
                }
            ]
        )

        self.assertEqual(
            build_roadmap_from_df(df),
            [
                (
                    "15/05/2026\n"
                    "Goal: Disponibilizar primeira versao\n"
                    "Marco: Publicacao do MVP",
                    10,
                ),
                (
                    "15/06/2026\n"
                    "Goal: Validar operacao\n"
                    "Marco: Encerrar piloto",
                    -10,
                ),
            ],
        )

    def test_legacy_selection_rejects_new_format_shape(self):
        path = self._save_workbook(self._new_format_workbook())
        try:
            with self.assertRaises(ValueError):
                load_base(path, planner_format="legacy")
        finally:
            os.unlink(path)

    def test_default_legacy_format_still_reads_first_task_sheet(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tarefas"
        ws.append(
            [
                "Nome da tarefa",
                "Nome do bucket",
                "Atribuído a",
                "Progresso",
                "Concluído em",
                "Data de início",
                "Data de conclusão",
                "Rótulos",
                "Notas",
            ]
        )
        ws.append(
            [
                "Implementar legado",
                "Concluído",
                "Dev",
                "Concluída",
                "2026-05-08",
                "2026-05-01",
                "2026-05-08",
                "HU002 - Legado [2SP];.DEV",
                "",
            ]
        )
        ws_plan = wb.create_sheet("Nome do plano")
        ws_plan.append(["Nome do plano", "Sprint Antiga"])
        ws_plan.append(["Exportado em", "2026-05-09"])

        path = self._save_workbook(wb)
        try:
            data = compute_json(path)
        finally:
            os.unlink(path)

        self.assertEqual(data["meta"]["sprint_name"], "Sprint Antiga")
        self.assertEqual(data["meta"]["export_date"], "2026-05-09")
        self.assertEqual(data["kpis"]["total"], 1)
        self.assertEqual(data["kpis"]["done"], 1)

    def test_teams_new_selection_falls_back_for_legacy_tarefas_sheet(self):
        path = self._save_workbook(self._legacy_format_workbook())
        try:
            data = compute_json(path, planner_format="teams_new")
        finally:
            os.unlink(path)

        self.assertEqual(data["meta"]["sprint_name"], "Sprint Antiga")
        self.assertEqual(data["meta"]["export_date"], "2026-05-09")
        self.assertEqual(data["meta"]["planner_format"], "legacy")
        self.assertEqual(data["kpis"]["total"], 1)
        self.assertEqual(data["kpis"]["done"], 1)

    def test_teams_new_selection_falls_back_for_legacy_first_sheet(self):
        path = self._save_workbook(self._legacy_format_workbook("Planner Export"))
        try:
            data = compute_json(path, planner_format="teams_new")
        finally:
            os.unlink(path)

        self.assertEqual(data["meta"]["sprint_name"], "Sprint Antiga")
        self.assertEqual(data["meta"]["export_date"], "2026-05-09")
        self.assertEqual(data["meta"]["planner_format"], "legacy")
        self.assertEqual(data["kpis"]["total"], 1)
        self.assertEqual(data["kpis"]["done"], 1)

    def test_replanned_new_hu_marker_enters_burndowns_on_marker_date(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tarefas"
        ws.append(
            [
                "Task name",
                "Bucket name",
                "Assigned to",
                "Status",
                "Date completed",
                "Start date",
                "Due date",
                "Created date",
                "Labels",
                "Notes",
            ]
        )
        ws.append(
            [
                "Implementar base",
                "Em Andamento",
                "Dev",
                "Em andamento",
                "",
                "2026-05-01",
                "2026-05-20",
                "2026-05-01",
                "HU001 - Base [5SP];.DEV",
                "",
            ]
        )
        ws.append(
            [
                "Implementar escopo novo",
                "Em Andamento",
                "Dev",
                "Em andamento",
                "",
                "2026-05-01",
                "2026-05-20",
                "2026-05-01",
                "HU002 - Escopo novo [3SP] [+E08/05];.DEV",
                "",
            ]
        )
        ws_plan = wb.create_sheet("Nome do plano")
        ws_plan.append(["Nome do plano", "Sprint Replanejada"])
        ws_plan.append(["Exportado em", "2026-05-14"])

        path = self._save_workbook(wb)
        try:
            data = compute_json(path)
        finally:
            os.unlink(path)

        self.assertEqual(data["burndown_sp"]["total_sp"], 8)
        self.assertEqual(data["burndown_sp"]["scope_events"][0]["hu"], "HU002")
        self.assertEqual(data["burndown_sp"]["scope_events"][0]["date"], "2026-05-08")
        self.assertEqual(data["burndown_sp"]["scope_events"][0]["sp"], 3.0)
        self.assertEqual(data["burndown_hu"]["scope_events"][0]["weight"], 1.0)

        sp_by_date = {row[0]: row for row in data["burndown_sp"]["rows"] if row[0]}
        self.assertEqual(sp_by_date["2026-05-07"][3], 5.0)
        self.assertEqual(sp_by_date["2026-05-07"][2], 5.0)
        self.assertEqual(sp_by_date["2026-05-08"][3], 8.0)
        self.assertEqual(sp_by_date["2026-05-08"][2], 8.0)
        self.assertEqual(sp_by_date["2026-05-08"][5], 3.0)

        hu_by_date = {row[0]: row for row in data["burndown_hu"]["rows"] if row[0]}
        self.assertEqual(hu_by_date["2026-05-07"][4], 1)
        self.assertEqual(hu_by_date["2026-05-07"][3], 1)
        self.assertEqual(hu_by_date["2026-05-08"][4], 2)
        self.assertEqual(hu_by_date["2026-05-08"][3], 2)
        self.assertEqual(hu_by_date["2026-05-08"][6], 1)

    def test_general_burndown_adds_nao_previsto_tasks_on_card_creation_date(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tarefas"
        ws.append(
            [
                "Task name",
                "Bucket name",
                "Assigned to",
                "Status",
                "Date completed",
                "Start date",
                "Due date",
                "Created date",
                "Labels",
                "Notes",
            ]
        )
        rows = [
            [
                "Tarefa planejada",
                "Em Andamento",
                "Dev",
                "Em andamento",
                "",
                "2026-05-01",
                "2026-05-20",
                "2026-05-01",
                ".DEV",
                "",
            ],
            [
                "Nao previsto aberto",
                "Em Andamento",
                "Dev",
                "Em andamento",
                "",
                "2026-05-08",
                "2026-05-20",
                "2026-05-08",
                "Nao Previsto;.DEV",
                "",
            ],
            [
                "Nao previsto concluido",
                "Concluido",
                "Dev",
                "Concluida",
                "2026-05-09",
                "2026-05-08",
                "2026-05-20",
                "2026-05-08",
                "Nao Previsto;.DEV",
                "",
            ],
        ]
        for row in rows:
            ws.append(row)

        ws_plan = wb.create_sheet("Nome do plano")
        ws_plan.append(["Nome do plano", "Sprint com Nao Previsto"])
        ws_plan.append(["Exportado em", "2026-05-14"])

        path = self._save_workbook(wb)
        try:
            data = compute_json(path)
        finally:
            os.unlink(path)

        self.assertEqual(data["kpis"]["total"], 3)
        self.assertEqual(data["kpis"]["nao_prev"], 2)
        self.assertEqual(data["burndown"]["scope_events"], [
            {"type": "add", "date": "2026-05-08", "label": "Nao Previsto", "count": 2}
        ])

        by_date = {row[0]: row for row in data["burndown"]["rows"] if row[0]}
        self.assertEqual(by_date["2026-05-07"][4], 1)
        self.assertEqual(by_date["2026-05-07"][3], 1)
        self.assertEqual(by_date["2026-05-07"][7], 0)
        self.assertEqual(by_date["2026-05-08"][4], 3)
        self.assertEqual(by_date["2026-05-08"][3], 3)
        self.assertEqual(by_date["2026-05-08"][6], 2)
        self.assertEqual(by_date["2026-05-09"][4], 3)
        self.assertEqual(by_date["2026-05-09"][3], 2)
        self.assertEqual(by_date["2026-05-09"][7], 1)

        self.assertEqual(data["burndown_nao_prev"]["scope_events"], [
            {
                "type": "add",
                "date": "2026-05-08",
                "hu": "",
                "label": "Fora de HU",
                "count": 2,
                "weight": 2,
            }
        ])
        nao_prev_by_date = {row[0]: row for row in data["burndown_nao_prev"]["rows"] if row[0]}
        self.assertEqual(nao_prev_by_date["2026-05-07"][4], 0)
        self.assertEqual(nao_prev_by_date["2026-05-07"][3], 0)
        self.assertEqual(nao_prev_by_date["2026-05-08"][4], 2)
        self.assertEqual(nao_prev_by_date["2026-05-08"][3], 2)
        self.assertEqual(nao_prev_by_date["2026-05-09"][4], 2)
        self.assertEqual(nao_prev_by_date["2026-05-09"][3], 1)
        self.assertEqual(nao_prev_by_date["2026-05-09"][7], 1)

    def test_nao_previsto_hu_burn_tracks_entries_and_completions(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tarefas"
        ws.append(
            [
                "Task name",
                "Bucket name",
                "Assigned to",
                "Status",
                "Date completed",
                "Start date",
                "Due date",
                "Created date",
                "Labels",
                "Notes",
            ]
        )
        rows = [
            [
                "Tarefa planejada",
                "Em Andamento",
                "Dev",
                "Em andamento",
                "",
                "2026-05-01",
                "2026-05-20",
                "2026-05-01",
                "HU001 - Base [5SP];.DEV",
                "",
            ],
            [
                "Nao previsto aberto HU010",
                "Em Andamento",
                "Dev",
                "Em andamento",
                "",
                "2026-05-08",
                "2026-05-20",
                "2026-05-08",
                "HU010 - Apoio [2SP];Nao Previsto;.DEV",
                "",
            ],
            [
                "Nao previsto concluido HU010",
                "Concluido",
                "Dev",
                "Concluida",
                "2026-05-09",
                "2026-05-08",
                "2026-05-20",
                "2026-05-08",
                "HU010 - Apoio [2SP];Nao Previsto;.DEV",
                "",
            ],
            [
                "Nao previsto aberto HU011",
                "Em Andamento",
                "Dev",
                "Em andamento",
                "",
                "2026-05-10",
                "2026-05-20",
                "2026-05-10",
                "HU011 - Ajuste [1SP];Nao Previsto;.DEV",
                "",
            ],
        ]
        for row in rows:
            ws.append(row)

        ws_plan = wb.create_sheet("Nome do plano")
        ws_plan.append(["Nome do plano", "Sprint com Nao Previsto em HU"])
        ws_plan.append(["Exportado em", "2026-05-14"])

        path = self._save_workbook(wb)
        try:
            data = compute_json(path)
        finally:
            os.unlink(path)

        scope_events = data["burndown_nao_prev_hu"]["scope_events"]
        self.assertEqual(scope_events[0]["date"], "2026-05-08")
        self.assertEqual(scope_events[0]["hu"], "HU010")
        self.assertEqual(scope_events[0]["count"], 2)
        self.assertEqual(scope_events[1]["date"], "2026-05-10")
        self.assertEqual(scope_events[1]["hu"], "HU011")
        self.assertEqual(scope_events[1]["count"], 1)

        by_date = {row[0]: row for row in data["burndown_nao_prev_hu"]["rows"] if row[0]}
        self.assertEqual(by_date["2026-05-07"][4], 0)
        self.assertEqual(by_date["2026-05-07"][3], 0)
        self.assertEqual(by_date["2026-05-08"][4], 2)
        self.assertEqual(by_date["2026-05-08"][3], 2)
        self.assertEqual(by_date["2026-05-09"][4], 2)
        self.assertEqual(by_date["2026-05-09"][3], 1)
        self.assertEqual(by_date["2026-05-09"][7], 1)
        self.assertEqual(by_date["2026-05-10"][4], 3)
        self.assertEqual(by_date["2026-05-10"][3], 2)


if __name__ == "__main__":
    unittest.main()
